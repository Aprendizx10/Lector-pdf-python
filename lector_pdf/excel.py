"""Exportación de resultados; no vuelve a leer los PDF."""
import json
import os
from pathlib import Path
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

def _valor(valor):
    if isinstance(valor,list): return ' / '.join(str(v) for v in valor)
    if isinstance(valor,dict): return json.dumps(valor,ensure_ascii=False)
    return valor

def _fila(hoja, valores):
    valores=[_valor(v) for v in valores]
    valores=[ILLEGAL_CHARACTERS_RE.sub('',v) if isinstance(v,str) else v for v in valores]
    hoja.append(valores)
    for i,valor in enumerate(valores,1):
        if isinstance(valor,str):
            c=hoja.cell(hoja.max_row,i)
            c.data_type='s'  # Conservar códigos y evitar interpretar fórmulas del PDF.
            c.number_format='@'

def exportar_excel(resultados, destino):
    if not resultados: raise ValueError('No hay resultados para exportar')
    destino=Path(destino)
    if destino.suffix.lower()!='.xlsx': raise ValueError('El archivo debe terminar en .xlsx')
    libro=Workbook()
    datos=libro.active
    datos.title='Datos'
    detalle=libro.create_sheet('Detalle')
    errores=libro.create_sheet('Errores')
    operaciones=libro.create_sheet('Comprobaciones')
    ids=list(dict.fromkeys(c['id'] for r in resultados for c in r.get('campos',[])))
    _fila(datos,['Archivo','ID documento','Estado del proceso']+ids)
    _fila(detalle,['Archivo','ID campo','Sección','Campo','Página','Valor original',
        'Original disponible','Valor normalizado','Estado de lectura','Validación',
        'Motivo','Método','Confianza OCR'])
    _fila(errores,['Archivo','Sección o campo','Error'])
    _fila(operaciones,['Archivo','Operación','Estado','Motivo','Total PDF','Total calculado','Diferencia'])
    for r in resultados:
        campos={c['id']:c.get('valor_normalizado') for c in r.get('campos',[])}
        _fila(datos,[r.get('archivo'),r.get('documento_id'),r.get('estado_proceso')]+[campos.get(k) for k in ids])
        for c in r.get('campos',[]):
            _fila(detalle,[r.get('archivo'),c.get('id'),c.get('seccion'),c.get('campo'),
                c.get('pagina'),c.get('valor_original'),c.get('original_disponible'),
                c.get('valor_normalizado'),c.get('lectura'),c.get('validacion'),
                c.get('motivo'),c.get('metodo'),c.get('confianza_ocr')])
        for e in r.get('errores',[]):
            _fila(errores,[r.get('archivo'),e.get('seccion') or e.get('campo'),e.get('mensaje')])
        comprobaciones=r.get('secciones',{}).get('financiera',{}).get('comprobaciones',{})
        for nombre,c in comprobaciones.items():
            _fila(operaciones,[r.get('archivo'),nombre,c.get('estado'),c.get('motivo'),
                c.get('valor_pdf'),c.get('valor_calculado'),c.get('diferencia')])
    for hoja in libro:
        hoja.freeze_panes='A2'
        hoja.auto_filter.ref=hoja.dimensions
        hoja.row_dimensions[1].height=42
        for c in hoja[1]:
            c.font=Font(bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor='17365D')
            c.alignment=Alignment(wrap_text=True,vertical='center')
        for i in range(1,hoja.max_column+1):
            hoja.column_dimensions[get_column_letter(i)].width=27
    # Guardar completo antes de sustituir un archivo existente.
    destino.parent.mkdir(parents=True,exist_ok=True)
    fd,temporal=tempfile.mkstemp(suffix='.xlsx',dir=destino.parent)
    os.close(fd)
    try:
        libro.save(temporal)
        os.replace(temporal,destino)
    finally:
        libro.close()
        if os.path.exists(temporal): os.unlink(temporal)
    return destino
