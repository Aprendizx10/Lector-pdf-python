import unittest
from tempfile import TemporaryDirectory
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from lector_pdf.excel import exportar_excel

class ExcelTests(unittest.TestCase):
    def datos(self):
        return [{'archivo':'ejemplo.pdf','documento_id':'abc','estado_proceso':'COMPLETADO',
          'campos':[{'id':'ciiu','campo':'CIIU','valor_normalizado':'0210','validacion':'OK'},
                    {'id':'importe','campo':'Importe','valor_normalizado':0,'validacion':'OK'},
                    {'id':'vacio','campo':'Vacío','valor_normalizado':None,'validacion':'OPCIONAL'},
                    {'id':'texto','campo':'Texto','valor_normalizado':'=1+1','validacion':'OK'}],
          'errores':[], 'secciones':{'financiera':{'comprobaciones':{
              'Total ingresos':{'estado':'NO VERIFICABLE','motivo':'Campo vacío','valor_pdf':100}}}}}]

    def test_valores_y_hojas(self):
        with TemporaryDirectory() as d:
            p=exportar_excel(self.datos(),Path(d)/'salida.xlsx')
            w=load_workbook(p)
            self.assertEqual(w.sheetnames,['Datos','Detalle','Errores','Comprobaciones'])
            h=w['Datos']
            self.assertEqual(h['D2'].value,'0210')
            self.assertEqual(h['D2'].data_type,'s')
            self.assertEqual(h['E2'].value,0)
            self.assertIsNone(h['F2'].value)
            self.assertEqual(h['G2'].value,'=1+1')
            self.assertEqual(h['G2'].data_type,'s')
            self.assertEqual(w['Comprobaciones']['C2'].value,'NO VERIFICABLE')
            self.assertEqual(h.freeze_panes,'A2')
            w.close()

    def test_error_se_conserva(self):
        with TemporaryDirectory() as d:
            p=exportar_excel([{'archivo':'roto.pdf','estado_proceso':'ERROR','errores':[{'mensaje':'Ilegible'}]}],Path(d)/'error.xlsx')
            w=load_workbook(p)
            self.assertEqual(w['Errores']['C2'].value,'Ilegible')
            w.close()

    def test_fallo_guardado_no_destruye_archivo(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'previo.xlsx'
            p.write_bytes(b'original')
            with patch('lector_pdf.excel.os.replace',side_effect=PermissionError('ocupado')):
                with self.assertRaises(PermissionError):exportar_excel(self.datos(),p)
            self.assertEqual(p.read_bytes(),b'original')
            self.assertEqual(len(list(Path(d).glob('*.xlsx'))),1)

    def test_sin_resultados(self):
        with self.assertRaises(ValueError):exportar_excel([],'salida.xlsx')

if __name__=='__main__':unittest.main()
